"""
Lógica del plano de siembra (sin interfaz, para poder probarla sola).

Un ENSAYO se describe con estas respuestas:
  tipo         -> tipo de ensayo (TEXTO LIBRE: SROW, MROW, o el que quieras)
  entradas     -> cuántas entradas tiene
  lineas       -> cuántas líneas tiene cada entrada
  ancho_lineas -> distancia entre líneas (m)
  tablones     -> cuántos tablones (rangos) tiene cada entrada
  cable        -> metros de cable usados para sembrar cada parcela (m)

De ahí salen las medidas reales:
  ancho de una entrada = lineas * ancho_lineas        (= ancho de la parcela)
  ancho total          = entradas * ancho de entrada
  alto total           = tablones * cable
  parcelas             = entradas * tablones
  área                 = ancho total * alto total  (m² y hectáreas)

Posición: (cx, cy) es el CENTRO del cuadro en metros sobre el lote
(0,0 = esquina inferior izquierda del lote).
Rotación: `angulo` en grados, sentido antihorario, 0 = tablones hacia arriba.

SENTIDO DE SIEMBRA (igual para todos los ensayos, y gira junto con ellos):
  - La PARCELA 1 (marcada con una X) es la de abajo a la izquierda (entrada 1, tablón 1).
  - Las ENTRADAS van de izquierda a derecha (flecha →, E1 -> En).
  - Los TABLONES van de abajo hacia arriba (flecha ↑, T1 -> Tn).
"""
import io
import math
import os
import re
import uuid

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from PIL import Image, ImageDraw, ImageFont
from shapely.affinity import rotate as _rotar
from shapely.geometry import Polygon, box
from shapely.ops import unary_union

from geo_kmz import nombre_rumbo, rumbo_entradas, rumbo_real

VERSION = "5.0"
TIPOS_SUGERIDOS = [
    "SROW", "MROW", "SPS", "INC DES", "GDM", "HILL", "PROGENY ROW",
    "YIELD TRIAL", "OBSERVACION", "VIVERO",
]
COLORES_SUGERIDOS = ["#4472C4", "#C44444", "#44A06A", "#C98A1B", "#8844C4", "#2A9CB8", "#B5651D", "#6AA121"]
TOL_SUPERPOSICION = 0.01  # m de penetración que se ignora (bordes que se tocan)
SENTIDO_TEXTO = "Entradas → de izquierda a derecha  ·  Tablones ↑ de abajo hacia arriba"
PARCELA1_TEXTO = "Abajo a la izquierda del ensayo (E1 · T1)"


def nuevo_id():
    return uuid.uuid4().hex[:8]


def tipos_conocidos(campos):
    """Tipos ya usados en cualquier campo + sugeridos (sin repetir, sin importar mayúsculas)."""
    out, vistos = [], set()
    usados = [e["tipo"] for cc in campos.values() for e in cc["ensayos"]]
    for t in usados + TIPOS_SUGERIDOS:
        t = str(t).strip()
        if t and t.casefold() not in vistos:
            vistos.add(t.casefold())
            out.append(t)
    return out


# ---------------------------------------------------------------- medidas
def dimensiones(e):
    ancho_entrada = e["lineas"] * e["ancho_lineas"]
    ancho = e["entradas"] * ancho_entrada
    alto = e["tablones"] * e["cable"]
    return {
        "ancho_entrada": round(ancho_entrada, 6),
        "ancho": round(ancho, 6),
        "alto": round(alto, 6),
        "parcelas": int(e["entradas"] * e["tablones"]),
        "area_m2": round(ancho * alto, 4),
        "area_ha": round(ancho * alto / 10000.0, 6),
        "parcela_m2": round(ancho_entrada * e["cable"], 4),
    }


def validar(e, otros_nombres=()):
    errores = []
    if not str(e.get("nombre", "")).strip():
        errores.append("Poné un nombre al ensayo.")
    elif e["nombre"].strip() in otros_nombres:
        errores.append(f"Ya hay otro ensayo llamado '{e['nombre'].strip()}'.")
    if not str(e.get("tipo", "")).strip():
        errores.append("Escribí el tipo de ensayo (el que quieras).")
    for campo, etiqueta in (("entradas", "entradas"), ("lineas", "líneas por entrada"), ("tablones", "tablones")):
        if int(e.get(campo) or 0) < 1:
            errores.append(f"La cantidad de {etiqueta} debe ser 1 o más.")
    if float(e.get("ancho_lineas") or 0) <= 0:
        errores.append("El ancho entre líneas debe ser mayor a 0.")
    if float(e.get("cable") or 0) <= 0:
        errores.append("El cable por parcela debe ser mayor a 0.")
    return errores


def numeracion_texto(e, max_items=6):
    """Qué números de parcela lleva cada entrada (cada entrada se numera de abajo hacia arriba)."""
    n_t, n_e = e["tablones"], e["entradas"]
    items = [f"E{i + 1}: {i * n_t + 1}–{(i + 1) * n_t}" for i in range(n_e)]
    if len(items) > max_items:
        items = items[: max_items - 1] + ["…"] + items[-1:]
    return "  ·  ".join(items) + "   (parcela 1 = X, abajo a la izquierda; cada entrada se numera de abajo hacia arriba)"


def rumbo_texto(lote, angulo):
    """Rumbo real hacia donde apuntan los tablones (↑), por ejemplo '358° (N)'; None sin KMZ."""
    r = rumbo_real(lote, angulo)
    return None if r is None else f"{r:.0f}° ({nombre_rumbo(r)})"


def rumbo_entradas_texto(lote, angulo):
    """Rumbo real hacia donde apuntan las entradas (→), por ejemplo '88° (E)'; None sin KMZ."""
    r = rumbo_entradas(lote, angulo)
    return None if r is None else f"{r:.0f}° ({nombre_rumbo(r)})"


# ---------------------------------------------------------------- geometría
def poligono(e):
    """Polígono real del cuadro (rotado alrededor de su centro)."""
    d = dimensiones(e)
    r = box(e["cx"] - d["ancho"] / 2, e["cy"] - d["alto"] / 2, e["cx"] + d["ancho"] / 2, e["cy"] + d["alto"] / 2)
    return _rotar(r, e.get("angulo", 0.0), origin=(e["cx"], e["cy"]))


def esquinas(e):
    return [(round(x, 3), round(y, 3)) for x, y in list(poligono(e).exterior.coords)[:4]]


def a_mundo(e, lx, ly):
    """Punto del ensayo (lx hacia la derecha, ly hacia arriba, desde su centro y SIN girar)
    -> metros del plano, ya con la rotación y la posición del ensayo."""
    t = math.radians(e.get("angulo", 0.0))
    c, s = math.cos(t), math.sin(t)
    return (e["cx"] + lx * c - ly * s, e["cy"] + lx * s + ly * c)


def _flecha_m(p0, p1, hs):
    """Polilínea de una flecha (eje + punta) entre p0 y p1; `hs` = largo de la punta."""
    dx, dy = p1[0] - p0[0], p1[1] - p0[1]
    L = math.hypot(dx, dy) or 1.0
    ux, uy = dx / L, dy / L
    nx, ny = -uy, ux
    bx, by = p1[0] - ux * hs, p1[1] - uy * hs
    return [p0, p1, (bx + nx * hs * 0.5, by + ny * hs * 0.5), p1, (bx - nx * hs * 0.5, by - ny * hs * 0.5)]


def marcas_ensayo(e, desvio=2.5):
    """Marcas de siembra de un ensayo, en metros del plano (listas de puntos x, y):
      parcela1          -> contorno de la parcela 1 (abajo a la izquierda)
      equis             -> las dos diagonales de la X que marca la parcela 1
      flecha_entradas   -> flecha → (E1 -> En), por debajo del ensayo
      flecha_tablones   -> flecha ↑ (T1 -> Tn), a la izquierda del ensayo
      centro_parcela1   -> punto central de la parcela 1
    """
    d = dimensiones(e)
    hw, hh = d["ancho"] / 2, d["alto"] / 2
    ae, cab = d["ancho_entrada"], float(e["cable"])
    largo_h, largo_v = max(2 * hw, 8.0), max(2 * hh, 8.0)
    hs_h = min(max(0.12 * largo_h, 1.2), 5.0)
    hs_v = min(max(0.12 * largo_v, 1.2), 5.0)
    M = lambda pts: [a_mundo(e, x, y) for x, y in pts]
    return {
        "parcela1": M([(-hw, -hh), (-hw + ae, -hh), (-hw + ae, -hh + cab), (-hw, -hh + cab)]),
        "equis": [M([(-hw, -hh), (-hw + ae, -hh + cab)]), M([(-hw + ae, -hh), (-hw, -hh + cab)])],
        "flecha_entradas": M(_flecha_m((-hw, -hh - desvio), (-hw + largo_h, -hh - desvio), hs_h)),
        "flecha_tablones": M(_flecha_m((-hw - desvio, -hh), (-hw - desvio, -hh + largo_v), hs_v)),
        "centro_parcela1": a_mundo(e, -hw + ae / 2, -hh + cab / 2),
    }


def _encogido(p):
    return p.buffer(-TOL_SUPERPOSICION, join_style=2)


def superposiciones(ensayos):
    """Lista de (nombre_a, nombre_b) de ensayos que se pisan."""
    polis = [_encogido(poligono(e)) for e in ensayos]
    out = []
    for i in range(len(ensayos)):
        for j in range(i + 1, len(ensayos)):
            if polis[i].intersects(polis[j]):
                out.append((ensayos[i]["nombre"], ensayos[j]["nombre"]))
    return out


def fuera_del_lote(e, poly):
    """True si el cuadro se sale del polígono del lote (sin lote: nunca)."""
    if poly is None:
        return False
    return not poly.buffer(TOL_SUPERPOSICION).contains(poligono(e))


def posicion_libre(existentes, ancho, alto, poly=None, paso=10.0):
    """Primer lugar libre (de abajo hacia arriba, de izquierda a derecha) donde
    entra un cuadro de ancho x alto sin pisar a otro y, si hay lote, sin salirse.
    Devuelve el CENTRO (cx, cy). Si no hay lugar, el centro del lote."""
    if poly is not None:
        minx, miny, maxx, maxy = poly.bounds
    else:
        minx, miny = 0.0, 0.0
        maxx = max([300.0] + [q[0] for e in existentes for q in esquinas(e)]) + ancho
        maxy = max([300.0] + [q[1] for e in existentes for q in esquinas(e)]) + alto
    otros = [_encogido(poligono(e)) for e in existentes]
    y = miny
    while y + alto <= maxy + 1e-9:
        x = minx
        while x + ancho <= maxx + 1e-9:
            cand = box(x, y, x + ancho, y + alto)
            if not any(cand.intersects(o) for o in otros) and (poly is None or poly.buffer(TOL_SUPERPOSICION).contains(cand)):
                return round(x + ancho / 2, 2), round(y + alto / 2, 2)
            x += paso
        y += paso
    if poly is not None:
        c = poly.centroid
        return round(c.x, 2), round(c.y, 2)
    return round(ancho / 2, 2), round(alto / 2, 2)


def choques_con_referencias(ensayos, lote, minimo_m2=1.0):
    """Ensayos que pisan un polígono de referencia del KMZ (por ejemplo, una siembra previa).
    Devuelve [(nombre_ensayo, nombre_referencia, m2_en_común)]."""
    if not lote:
        return []
    zonas = []
    for r in lote.get("refs", []):
        if r["k"] == "poly" and len(r["c"]) >= 3:
            p = Polygon(r["c"])
            zonas.append((r["n"] or "polígono sin nombre", p if p.is_valid else p.buffer(0)))
    out = []
    for e in ensayos:
        pe = poligono(e)
        for nombre, z in zonas:
            area = pe.intersection(z).area
            if area >= minimo_m2:
                out.append((e["nombre"], nombre, area))
    return out


# ---------------------------------------------------------------- áreas
def ocupacion(ensayos, poly):
    """Cuánto del lote ocupan los ensayos (lo superpuesto o lo que queda fuera
    del lote no se cuenta dos veces / no se cuenta)."""
    lote_ha = poly.area / 10000.0 if poly is not None else None
    if not ensayos:
        return {"suma_ha": 0.0, "ocupado_ha": 0.0, "lote_ha": lote_ha,
                "libre_ha": lote_ha, "pct_ocupado": 0.0 if lote_ha else None}
    suma = sum(dimensiones(e)["area_ha"] for e in ensayos)
    union = unary_union([poligono(e) for e in ensayos])
    ocupado = (union.intersection(poly).area if poly is not None else union.area) / 10000.0
    return {
        "suma_ha": suma, "ocupado_ha": ocupado, "lote_ha": lote_ha,
        "libre_ha": (lote_ha - ocupado) if lote_ha is not None else None,
        "pct_ocupado": (ocupado / lote_ha * 100.0) if lote_ha else None,
    }


def pct_del_lote(e, poly):
    if poly is None:
        return None
    return dimensiones(e)["area_m2"] / poly.area * 100.0


# ---------------------------------------------------------------- dibujos (imágenes)
_DIR_FUENTES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fuentes")


def _cargar_fuente(nombres, tam):
    """Prueba cada fuente en orden. La primera es la que viene dentro de la carpeta `fuentes`, así los
    dibujos (flechas →, ↑ y letras) salen igual en cualquier computadora o servidor."""
    for nombre in nombres:
        try:
            return ImageFont.truetype(nombre, tam)
        except OSError:
            pass
    try:
        return ImageFont.load_default(size=tam)
    except TypeError:
        return ImageFont.load_default()


def _fuente(tam):
    return _cargar_fuente((os.path.join(_DIR_FUENTES, "DejaVuSans.ttf"), "DejaVuSans.ttf", "arial.ttf",
                           "LiberationSans-Regular.ttf"), tam)


def _fuente_negrita(tam):
    return _cargar_fuente((os.path.join(_DIR_FUENTES, "DejaVuSans-Bold.ttf"), "DejaVuSans-Bold.ttf", "arialbd.ttf",
                           "LiberationSans-Bold.ttf"), tam)


def _rgb(hexcolor):
    h = hexcolor.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def _largo_bonito(max_m):
    mejor = 1.0
    for p in range(-1, 6):
        for b in (1, 2, 5):
            v = b * 10 ** p
            if v <= max_m:
                mejor = v
    return mejor


def _png(im):
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


NEGRO = (17, 17, 17, 255)
HALO = (255, 255, 255, 240)


def _flecha_img(dr, p0, p1, grosor=3, cabeza=11.0):
    """Flecha negra con halo blanco entre dos puntos (px de imagen)."""
    dx, dy = p1[0] - p0[0], p1[1] - p0[1]
    L = math.hypot(dx, dy) or 1.0
    ux, uy = dx / L, dy / L
    nx, ny = -uy, ux
    h = min(cabeza, max(6.0, L * 0.25))
    b = (p1[0] - ux * h, p1[1] - uy * h)
    punta = [p1, (b[0] + nx * h * 0.55, b[1] + ny * h * 0.55), (b[0] - nx * h * 0.55, b[1] - ny * h * 0.55)]
    dr.line([p0, b], fill=HALO, width=grosor + 3)
    dr.polygon(punta, fill=HALO, outline=HALO, width=3)
    dr.line([p0, b], fill=NEGRO, width=grosor)
    dr.polygon(punta, fill=NEGRO)


def _equis_img(dr, c4, grosor=3):
    """X negra con halo blanco que une las esquinas opuestas de un cuadro (c4 = 4 esquinas en orden)."""
    for a, b in ((c4[0], c4[2]), (c4[1], c4[3])):
        dr.line([a, b], fill=HALO, width=grosor + 3)
    for a, b in ((c4[0], c4[2]), (c4[1], c4[3])):
        dr.line([a, b], fill=NEGRO, width=grosor)


def mapa_png(ensayos, poly, titulo="", ancho_px=1100, alto_max_px=760):
    """Mapa a escala del campo: lote, ensayos (con su rotación), la X de la parcela 1, las
    flechas de sentido de siembra, escala y norte."""
    polis = [(e, poligono(e)) for e in ensayos]
    xs, ys = [], []
    if poly is not None:
        xs += [p[0] for p in poly.exterior.coords]
        ys += [p[1] for p in poly.exterior.coords]
    for _, p in polis:
        xs += [q[0] for q in p.exterior.coords]
        ys += [q[1] for q in p.exterior.coords]
    if not xs:
        xs, ys = [0, 100], [0, 100]
    minx, maxx, miny, maxy = min(xs), max(xs), min(ys), max(ys)
    bw, bh = max(maxx - minx, 1.0), max(maxy - miny, 1.0)
    pad, pad_sup = 46, 70          # más aire arriba para el título y la leyenda
    esc = min((ancho_px - 2 * pad) / bw, (alto_max_px - pad - pad_sup) / bh)
    W, H = max(int(bw * esc + 2 * pad), 520), max(int(bh * esc + pad + pad_sup), 300)
    tx = lambda x: pad + (x - minx) * esc
    ty = lambda y: H - pad - (y - miny) * esc
    P = lambda q: (tx(q[0]), ty(q[1]))

    im = Image.new("RGB", (W, H), (250, 250, 250))
    d = ImageDraw.Draw(im, "RGBA")
    f_chica, f_media, f_gran = _fuente(13), _fuente(15), _fuente(18)

    if poly is not None:
        pts = [(tx(x), ty(y)) for x, y in poly.exterior.coords]
        d.polygon(pts, fill=(21, 101, 192, 36))
        d.line(pts + [pts[0]], fill=(21, 101, 192, 255), width=3)

    pisados = {n for par in superposiciones(ensayos) for n in par}
    etiquetas = []
    for e, p in polis:
        pts = [(tx(x), ty(y)) for x, y in list(p.exterior.coords)[:4]]
        col = _rgb(e["color"])
        malo = e["nombre"] in pisados
        afuera = fuera_del_lote(e, poly)
        borde = (211, 47, 47, 255) if malo else ((239, 108, 0, 255) if afuera else (38, 50, 55, 255))
        d.polygon(pts, fill=col + ((150,) if (malo or afuera) else (225,)))
        d.line(pts + [pts[0]], fill=borde, width=3 if (malo or afuera) else 2)
        dim = dimensiones(e)
        # X en la parcela 1 y flechas de sentido (en píxeles, con largos mínimos para que se vean a esta escala)
        t = math.radians(e.get("angulo", 0.0))
        up, right = (-math.sin(t), math.cos(t)), (math.cos(t), math.sin(t))
        w, h = dim["ancho"], dim["alto"]
        esq = (e["cx"] - right[0] * w / 2 - up[0] * h / 2, e["cy"] - right[1] * w / 2 - up[1] * h / 2)
        c1 = (esq[0] + right[0] * dim["ancho_entrada"] / 2 + up[0] * e["cable"] / 2,
              esq[1] + right[1] * dim["ancho_entrada"] / 2 + up[1] * e["cable"] / 2)
        cx1, cy1 = P(c1)
        r = 5
        _equis_img(d, [(cx1 - r, cy1 - r), (cx1 + r, cy1 - r), (cx1 + r, cy1 + r), (cx1 - r, cy1 + r)], grosor=2)
        o = 7 / esc
        e0 = (esq[0] - up[0] * o, esq[1] - up[1] * o)
        le = max(w, 26 / esc)
        _flecha_img(d, P(e0), P((e0[0] + right[0] * le, e0[1] + right[1] * le)), grosor=2, cabeza=8)
        v0 = (esq[0] - right[0] * o, esq[1] - right[1] * o)
        lv = max(h, 26 / esc)
        _flecha_img(d, P(v0), P((v0[0] + up[0] * lv, v0[1] + up[1] * lv)), grosor=2, cabeza=8)
        etiquetas.append((e, dim, pts))

    # etiquetas: sin pisarse entre sí y sin salirse de la imagen
    ocupadas = []
    for e, dim, pts in etiquetas:
        txt = f"{e['nombre']}\n{dim['area_ha']:.3f} ha"
        bb = d.multiline_textbbox((0, 0), txt, font=f_media, spacing=2)
        w, h = bb[2] + 8, bb[3] + 6
        bx0, bx1 = min(q[0] for q in pts), max(q[0] for q in pts)
        cy = sum(q[1] for q in pts) / 4
        x = bx1 + 5 if bx1 + 5 + w <= W - 4 else max(4, bx0 - 5 - w)
        y = min(max(cy - h / 2, 4), H - h - 4)
        for _ in range(40):
            choque = next((o for o in ocupadas if not (x + w <= o[0] or o[2] <= x or y + h <= o[1] or o[3] <= y)), None)
            if choque is None:
                break
            y = choque[3] + 2
            if y + h > H - 4:
                y = 4
                x = max(4, x - w - 4)
        ocupadas.append((x, y, x + w, y + h))
        d.rectangle([x, y, x + w, y + h], fill=(255, 255, 255, 225), outline=(176, 190, 197, 255))
        d.multiline_text((x + 4, y + 3), txt, fill=(38, 50, 55, 255), font=f_media, spacing=2)

    # escala
    largo = _largo_bonito(120 / esc)
    x0, y0 = 16, H - 18
    d.line([(x0, y0), (x0 + largo * esc, y0)], fill=(55, 71, 79, 255), width=4)
    d.line([(x0, y0 - 5), (x0, y0 + 5)], fill=(55, 71, 79, 255), width=2)
    d.line([(x0 + largo * esc, y0 - 5), (x0 + largo * esc, y0 + 5)], fill=(55, 71, 79, 255), width=2)
    d.text((x0, y0 - 24), f"{largo:g} m", fill=(55, 71, 79, 255), font=f_chica)
    # norte
    nx, ny = W - 28, 52
    d.polygon([(nx, ny - 26), (nx - 9, ny + 6), (nx, ny), (nx + 9, ny + 6)], fill=(55, 71, 79, 255))
    d.text((nx - 5, ny + 10), "N", fill=(55, 71, 79, 255), font=f_media)
    if titulo:
        d.text((14, 12), titulo, fill=(33, 33, 33, 255), font=f_gran)
    d.text((14, 36), "X = parcela 1  ·  → entradas (E1 a En)  ·  ↑ tablones (T1 a Tn)", fill=(55, 71, 79, 255), font=f_chica)
    return _png(im)


def escala_comun(ensayos, max_px=560.0, tope=14.0):
    """px por metro para dibujar TODAS las fotos de un campo a la misma escala
    (así se pueden comparar tamaños entre ensayos)."""
    ext = 0.0
    for e in ensayos:
        q = esquinas({**e, "cx": 0.0, "cy": 0.0})
        ext = max(ext, max(p[0] for p in q) - min(p[0] for p in q), max(p[1] for p in q) - min(p[1] for p in q))
    return min(tope, max_px / ext) if ext > 0 else tope


def ficha_png(e, escala, pct=None, rumbo_tab=None, rumbo_ent=None):
    """Foto de UN ensayo tal cual está en el plano: con su rotación, las entradas (columnas) y los
    tablones, la X en la parcela 1, la flecha → de las entradas (E1 a En) y la flecha ↑ de los
    tablones (T1 a Tn), la escala y los datos principales."""
    d = dimensiones(e)
    ang = float(e.get("angulo", 0.0))
    t = math.radians(ang)
    c, s = math.cos(t), math.sin(t)
    hw, hh = max(d["ancho"] * escala / 2, 2.0), max(d["alto"] * escala / 2, 2.0)
    n_e, n_t = e["entradas"], e["tablones"]
    we, ht = 2 * hw / n_e, 2 * hh / n_t

    def off(lx, ly):   # coordenadas del ensayo (x der., ly = "arriba") -> imagen (x der., y abajo)
        return (lx * c - ly * s, -(lx * s + ly * c))

    esq = [off(-hw, -hh), off(hw, -hh), off(hw, hh), off(-hw, hh)]
    xw, xh = max(we, 12.0), max(ht, 12.0)                       # la X ocupa la parcela 1 (mínimo 12 px)
    cuadro_x = [off(-hw, -hh), off(-hw + xw, -hh), off(-hw + xw, -hh + xh), off(-hw, -hh + xh)]
    largo_h, largo_v = max(2 * hw, 44.0), max(2 * hh, 44.0)
    ah0, ah1 = off(-hw, -hh - 14), off(-hw + largo_h, -hh - 14)    # flecha de las entradas (→), abajo
    av0, av1 = off(-hw - 14, -hh), off(-hw - 14, -hh + largo_v)    # flecha de los tablones (↑), a la izquierda
    rotulos = [("E1", off(-hw + 8, -hh - 31)), ("T1", off(-hw - 33, -hh + 8))]
    if n_e > 1:
        rotulos.append((f"E{n_e}", off(-hw + largo_h - 8, -hh - 31)))
    if n_t > 1:
        rotulos.append((f"T{n_t}", off(-hw - 34, -hh + largo_v - 8)))
    todos = esq + cuadro_x + [ah0, ah1, av0, av1] + [p for _, p in rotulos]
    minx, maxx = min(p[0] for p in todos) - 26, max(p[0] for p in todos) + 26
    miny, maxy = min(p[1] for p in todos) - 16, max(p[1] for p in todos) + 16
    x_txt = (maxx - minx) + 6
    W = int(x_txt + 300)
    H = max(int(maxy - miny) + 36, 300)
    Pq = lambda p: (p[0] - minx, p[1] - miny + 8)

    im = Image.new("RGB", (W, H), (255, 255, 255))
    dr = ImageDraw.Draw(im, "RGBA")
    f_chica, f_media = _fuente(13), _fuente(14)
    f_gran, f_ha = _fuente_negrita(18), _fuente_negrita(22)
    col = _rgb(e["color"])

    dr.polygon([Pq(q) for q in esq], fill=col + (230,))
    if we >= 3:
        for k in range(1, n_e):
            dr.line([Pq(off(-hw + k * we, -hh)), Pq(off(-hw + k * we, hh))], fill=(255, 255, 255, 200), width=1)
    if ht >= 4:
        for k in range(1, n_t):
            dr.line([Pq(off(-hw, -hh + k * ht)), Pq(off(hw, -hh + k * ht))], fill=(255, 255, 255, 140), width=1)
    pts = [Pq(q) for q in esq]
    dr.line(pts + [pts[0]], fill=(38, 50, 55, 255), width=2)
    _equis_img(dr, [Pq(q) for q in cuadro_x], grosor=3)          # X = parcela 1
    _flecha_img(dr, Pq(ah0), Pq(ah1), grosor=3, cabeza=12)       # entradas →
    _flecha_img(dr, Pq(av0), Pq(av1), grosor=3, cabeza=12)       # tablones ↑
    for texto, p in rotulos:
        dr.text(Pq(p), texto, fill=(55, 71, 79, 255), font=f_chica, anchor="mm")

    # datos a la derecha
    y = 12
    lineas = [
        (e["nombre"], f_gran, (33, 33, 33, 255)),
        (str(e["tipo"]), f_media, (84, 110, 122, 255)),
        (f"{d['area_ha']:.3f} ha", f_ha, col + (255,)),
    ]
    if pct is not None:
        lineas.append((f"{pct:.1f}% del lote", f_media, (55, 71, 79, 255)))
    lineas += [
        (f"{d['ancho']:g} × {d['alto']:g} m", f_media, (55, 71, 79, 255)),
        (f"{n_e} entradas × {n_t} tablones", f_media, (55, 71, 79, 255)),
        (f"rotación {ang:g}°", f_media, (55, 71, 79, 255)),
        ("X   parcela 1 (E1 · T1)", f_media, (17, 17, 17, 255)),
        (f"→   entradas E1 a E{n_e}" + (f" · {rumbo_ent}" if rumbo_ent else ""), f_media, (17, 17, 17, 255)),
        (f"↑   tablones T1 a T{n_t}" + (f" · {rumbo_tab}" if rumbo_tab else ""), f_media, (17, 17, 17, 255)),
    ]
    for txt, f, colr in lineas:
        dr.text((x_txt, y), txt, fill=colr, font=f)
        y += f.size + 7
    # escala
    largo = _largo_bonito(90 / escala)
    x0, y0 = 14, H - 16
    dr.line([(x0, y0), (x0 + largo * escala, y0)], fill=(55, 71, 79, 255), width=4)
    dr.text((x0, y0 - 20), f"{largo:g} m", fill=(55, 71, 79, 255), font=f_chica)
    dr.rectangle([0, 0, W - 1, H - 1], outline=(207, 216, 220, 255))
    return _png(im), W, H


# ---------------------------------------------------------------- Excel
BORDE = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
PX_FILA = 20  # una fila de 15 pt = 20 px


def _bandas(ensayos):
    """Ordena los ensayos de norte a sur y, dentro de una misma franja del
    campo, de izquierda a derecha."""
    items = []
    for e in ensayos:
        ys = [q[1] for q in esquinas(e)]
        items.append((e, min(ys), max(ys)))
    items.sort(key=lambda t: -t[2])
    bandas = []
    for e, y0, y1 in items:
        h = y1 - y0
        for b in bandas:
            solape = min(y1, b["y1"]) - max(y0, b["y0"])
            if solape > 0.5 * min(h, b["y1"] - b["y0"]):
                b["items"].append(e)
                b["y0"], b["y1"] = min(b["y0"], y0), max(b["y1"], y1)
                break
        else:
            bandas.append({"y0": y0, "y1": y1, "items": [e]})
    bandas.sort(key=lambda b: -b["y1"])
    for b in bandas:
        b["items"].sort(key=lambda e: e["cx"])
    return bandas


def _nombre_hoja(nombre, usados):
    base = re.sub(r"[\[\]:*?/\\]", "-", nombre).strip()[:31] or "Campo"
    final, n = base, 2
    while final in usados:
        suf = f" ({n})"
        final = base[: 31 - len(suf)] + suf
        n += 1
    usados.add(final)
    return final


def _pegar_imagen(ws, png, celda):
    img = XLImage(io.BytesIO(png))
    ws.add_image(img, celda)
    return img


def _ficha(ws, e, fila, escala, lote):
    """Una ficha: tabla de datos a la izquierda y la foto del ensayo a la derecha.
    Devuelve la próxima fila libre."""
    poly = lote["poly"] if lote else None
    d = dimensiones(e)
    pct = pct_del_lote(e, poly)
    ang = e.get("angulo", 0.0)
    rt, re_ = rumbo_texto(lote, ang), rumbo_entradas_texto(lote, ang)
    color = e["color"].lstrip("#")
    png, w_px, h_px = ficha_png(e, escala, pct, rt, re_)
    _pegar_imagen(ws, png, f"D{fila}")

    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=2)
    t = ws.cell(row=fila, column=1, value=f"{e['nombre']}  ·  {e['tipo']}  ·  {d['area_ha']:.3f} ha")
    t.font = Font(bold=True, color="FFFFFF", size=12)
    t.alignment = Alignment(vertical="center")
    for col in (1, 2):
        ws.cell(row=fila, column=col).fill = PatternFill("solid", fgColor=color)
    ws.row_dimensions[fila].height = 22

    filas = [
        ("Tipo de ensayo", e["tipo"], None),
        ("HECTÁREAS del ensayo (ha)", d["area_ha"], "0.000"),
        ("% del lote", (pct / 100.0) if pct is not None else "sin KMZ", "0.0%"),
        ("Área (m²)", d["area_m2"], "#,##0"),
        ("Entradas", e["entradas"], None),
        ("Líneas por entrada", e["lineas"], None),
        ("Ancho entre líneas (m)", e["ancho_lineas"], None),
        ("Ancho de cada entrada (m)", d["ancho_entrada"], None),
        ("Tablones (rangos) por entrada", e["tablones"], None),
        ("Cable por parcela (m)", e["cable"], None),
        ("Tamaño de cada parcela (m)", f"{d['ancho_entrada']:g} × {e['cable']:g}", None),
        ("Parcelas totales", d["parcelas"], "#,##0"),
        ("Ancho × alto del ensayo (m)", f"{d['ancho']:g} × {d['alto']:g}", None),
        ("Parcela 1 (marcada con X)", PARCELA1_TEXTO, None),
        ("Sentido de siembra", SENTIDO_TEXTO, None),
        ("Rotación (°, antihorario)", ang, "0.0"),
        ("Rumbo de los tablones ↑ (T1 a Tn)", rt or "sin KMZ", None),
        ("Rumbo de las entradas → (E1 a En)", re_ or "sin KMZ", None),
        ("Centro en el lote X ; Y (m)", f"{e['cx']:.1f} ; {e['cy']:.1f}", None),
        ("Numeración de parcelas", numeracion_texto(e), None),
    ]
    altos = {"Parcela 1 (marcada con X)": 30, "Sentido de siembra": 30, "Numeración de parcelas": 60}
    for i, (rot, val, fmt) in enumerate(filas, 1):
        a = ws.cell(row=fila + i, column=1, value=rot)
        b = ws.cell(row=fila + i, column=2, value=val)
        for c in (a, b):
            c.border = BORDE
            c.alignment = Alignment(vertical="center", wrap_text=True, horizontal="left")
        a.font = Font(size=10, color="37474F")
        a.fill = PatternFill("solid", fgColor="ECEFF1")
        if fmt:
            b.number_format = fmt
        if rot.startswith("HECTÁREAS"):
            a.font = Font(bold=True, size=11)
            b.font = Font(bold=True, size=12, color=color)
            a.fill = PatternFill("solid", fgColor="FFF8E1")
            b.fill = PatternFill("solid", fgColor="FFF8E1")
        ws.row_dimensions[fila + i].height = altos.get(rot, 15)
    filas_tabla = 1 + len(filas) + 2
    filas_foto = math.ceil(h_px / PX_FILA) + 1
    for r in range(fila + len(filas) + 1, fila + max(filas_tabla, filas_foto) + 1):
        ws.row_dimensions[r].height = 15
    return fila + max(filas_tabla, filas_foto) + 2


def _escribir_campo(ws, campo):
    ensayos, lote = campo["ensayos"], campo.get("lote")
    poly = lote["poly"] if lote else None
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 3
    ws.cell(row=1, column=1, value=f"PLANO DE SIEMBRA — {campo['nombre']}").font = Font(bold=True, size=14)
    oc = ocupacion(ensayos, poly)
    resumen = f"Ensayos: {len(ensayos)}  ·  Hectáreas de los ensayos: {oc['suma_ha']:.3f} ha"
    if oc["lote_ha"] is not None:
        resumen += (f"  ·  Lote: {oc['lote_ha']:.2f} ha  ·  Ocupado: {oc['ocupado_ha']:.3f} ha "
                    f"({oc['pct_ocupado']:.1f}%)  ·  Libre: {oc['libre_ha']:.2f} ha")
    ws.cell(row=2, column=1, value=resumen).font = Font(bold=True, size=11, color="37474F")
    ws.cell(row=3, column=1, value="Mapa a escala (norte arriba). X = parcela 1 · → entradas (E1 a En) · ↑ tablones (T1 a Tn). "
                                   "Abajo, una ficha con foto por ensayo: las fotos están a la misma escala y son imágenes que "
                                   "podés mover, achicar o reemplazar.").font = Font(italic=True, size=9, color="555555")

    mapa = mapa_png(ensayos, poly, titulo=campo["nombre"])
    img = _pegar_imagen(ws, mapa, "A5")
    filas_mapa = math.ceil(img.height / PX_FILA) + 1
    for r in range(4, 5 + filas_mapa):
        ws.row_dimensions[r].height = 15
    fila = 5 + filas_mapa + 1
    ws.row_breaks.append(Break(id=fila - 1))  # al imprimir, el mapa va solo en su página

    esc = escala_comun(ensayos)
    for banda in _bandas(ensayos):
        for e in banda["items"]:
            fila = _ficha(ws, e, fila, esc, lote)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _tabla(ws, fila, encabezados, filas, formatos=None, negrita_ultima=False):
    for i, h in enumerate(encabezados, 1):
        c = ws.cell(row=fila, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="37474F")
        c.border = BORDE
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[fila].height = 44
    for r, valores in enumerate(filas, fila + 1):
        for i, v in enumerate(valores, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BORDE
            if formatos and formatos.get(i):
                c.number_format = formatos[i]
            if negrita_ultima and r == fila + len(filas):
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="ECEFF1")
    return fila + len(filas) + 1


def generar_workbook(campos):
    """campos: lista de {"nombre": str, "ensayos": [ensayo, ...], "lote": dict de geo_kmz | None}"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.cell(row=1, column=1, value="RESUMEN DE ENSAYOS Y HECTÁREAS").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value="Sentido de siembra en todos los ensayos: parcela 1 = X (abajo a la izquierda) · "
                                   "entradas → de izquierda a derecha · tablones ↑ de abajo hacia arriba. "
                                   "Todo gira junto con el ensayo.").font = Font(italic=True, size=9, color="555555")

    ws.cell(row=3, column=1, value="Por campo").font = Font(bold=True, size=11)
    filas_campo = []
    for c in campos:
        oc = ocupacion(c["ensayos"], c["lote"]["poly"] if c.get("lote") else None)
        filas_campo.append([
            c["nombre"], oc["lote_ha"], len(c["ensayos"]), sum(dimensiones(e)["parcelas"] for e in c["ensayos"]),
            oc["suma_ha"], oc["ocupado_ha"] if oc["lote_ha"] is not None else None,
            (oc["pct_ocupado"] / 100.0) if oc["pct_ocupado"] is not None else None, oc["libre_ha"],
        ])
    sig = _tabla(ws, 4, ["Campo", "Área del lote (ha)", "Ensayos", "Parcelas", "Hectáreas de los ensayos (ha)",
                         "Ocupado dentro del lote (ha)", "% del lote ocupado", "Libre (ha)"],
                 filas_campo, formatos={2: "0.00", 4: "#,##0", 5: "0.000", 6: "0.000", 7: "0.0%", 8: "0.00"})

    ws.cell(row=sig + 1, column=1, value="Por ensayo").font = Font(bold=True, size=11)
    filas_ens, tot_ha, tot_parc = [], 0.0, 0
    for c in campos:
        poly = c["lote"]["poly"] if c.get("lote") else None
        for e in c["ensayos"]:
            d = dimensiones(e)
            pct = pct_del_lote(e, poly)
            ang = e.get("angulo", 0.0)
            tot_ha += d["area_ha"]
            tot_parc += d["parcelas"]
            filas_ens.append([
                c["nombre"], e["nombre"], e["tipo"], d["area_ha"], (pct / 100.0) if pct is not None else None,
                round(ang, 1), rumbo_texto(c.get("lote"), ang), rumbo_entradas_texto(c.get("lote"), ang),
                e["entradas"], e["lineas"], e["ancho_lineas"], e["cable"],
                e["tablones"], d["parcelas"], f"{d['ancho_entrada']:g} × {e['cable']:g}", d["ancho"], d["alto"],
                round(e["cx"], 1), round(e["cy"], 1),
            ])
    filas_ens.append(["TOTAL", None, None, tot_ha] + [None] * 9 + [tot_parc] + [None] * 5)
    _tabla(ws, sig + 2, ["Campo", "Ensayo", "Tipo", "HECTÁREAS (ha)", "% del lote", "Rotación (°)",
                         "Rumbo de los tablones ↑", "Rumbo de las entradas →", "Entradas", "Líneas por entrada",
                         "Ancho entre líneas (m)", "Cable por parcela (m)", "Tablones por entrada", "Parcelas",
                         "Tamaño de parcela (m)", "Ancho del ensayo (m)", "Alto del ensayo (m)", "Centro X (m)",
                         "Centro Y (m)"], filas_ens,
           formatos={4: "0.000", 5: "0.0%", 14: "#,##0"}, negrita_ultima=True)
    for i in range(1, 20):
        ws.column_dimensions[get_column_letter(i)].width = 15
    ws.column_dimensions["A"].width = 20
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    usados = {"Resumen"}
    for c in campos:
        hoja = wb.create_sheet(_nombre_hoja(c["nombre"], usados))
        _escribir_campo(hoja, c)
    wb.active = 0
    return wb
